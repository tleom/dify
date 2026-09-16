"""Planning container policy plus opt-in tests against a real Docker engine.

Set WORKBENCH_PLAN_TEST_IMAGE to a locally available sandbox/Python image to run
the Docker cases. Fixtures use only unique test-owned containers and volumes.
"""

import importlib.util
import json
import os
from pathlib import Path
import subprocess
import threading
import time
from types import SimpleNamespace
from uuid import uuid4

import pytest


@pytest.fixture
def planning():
    path = Path(__file__).parents[1] / "sandbox-manager/planning.py"
    spec = importlib.util.spec_from_file_location("planning_manager_under_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def payload(**changes):
    return {
        "binding_id": str(uuid4()),
        "execution_id": str(uuid4()),
        "request_key": "inspection",
        "script": "printf read",
        "timeout": 2,
        **changes,
    }


@pytest.fixture
def admission(planning, tmp_path):
    owner_lock = threading.RLock()
    return planning.InspectionAdmission(tmp_path, lambda key: owner_lock)


def test_container_launch_uses_readonly_owner_mounts_and_no_runtime_credentials(
    planning, admission, monkeypatch
):
    calls = []

    def docker(*args, **kwargs):
        calls.append((args, kwargs))
        return SimpleNamespace(
            returncode=0, stdout="0" if args[0] == "wait" else '{"output": "read"}'
        )

    monkeypatch.setattr(planning, "copy_global_resources", lambda *args: None)
    request = payload()
    request["admission_ticket"] = admission.ticket("owner", request["binding_id"])
    result = planning.inspect_plan(
        "owner",
        request,
        docker=docker,
        ensure=lambda key: None,
        identity=lambda key: ("test-owner", "must-not-propagate"),
        image="test-image",
        prefix="test",
        manager_python=("/usr/local/bin/python", "-I", "-S"),
        admission=admission,
    )
    command = next(
        args for args, _ in calls if args[0] == "create" and "--volumes-from" in args
    )
    assert command[command.index("--volumes-from") + 1] == "test-owner:ro"
    assert command[command.index("--network") + 1] == "none"
    assert command[command.index("--user") + 1] == "1000"
    assert "--read-only" in command and "--cap-add" not in command
    assert command[command.index("--cap-drop") + 1] == "ALL"
    assert "must-not-propagate" not in " ".join(command)
    assert calls[-1][0][:2] == ("rm", "-f") and result == {"output": "read"}


def test_duplicate_create_never_removes_an_existing_investigation(
    planning, admission, monkeypatch
):
    removed = []

    def docker(*args, **kwargs):
        if args[0] == "create":
            raise RuntimeError("container already exists")
        if args[0] == "rm":
            removed.append(args)
        return SimpleNamespace(returncode=0, stdout="")

    request = payload()
    request["admission_ticket"] = admission.ticket("owner", request["binding_id"])
    with pytest.raises(RuntimeError, match="already exists"):
        planning.inspect_plan(
            "owner",
            request,
            docker=docker,
            ensure=lambda key: None,
            identity=lambda key: ("test-owner", "token"),
            image="test-image",
            prefix="test",
            manager_python=("/usr/local/bin/python", "-I", "-S"),
            admission=admission,
        )
    assert removed == []


@pytest.fixture
def real_inspection(planning, admission):
    image = os.environ.get("WORKBENCH_PLAN_TEST_IMAGE")
    if not image:
        pytest.skip("Set WORKBENCH_PLAN_TEST_IMAGE for real Docker isolation tests")
    prefix = "wb-plan-test-" + uuid4().hex[:12]
    owner, key, binding = prefix + "-owner", str(uuid4()), str(uuid4())

    def docker(*args, stdin=None, timeout=90, check=True):
        return subprocess.run(
            ["docker", *args],
            input=stdin,
            capture_output=True,
            text=True,
            timeout=timeout,
            check=check,
            encoding="utf-8",
        )

    docker("image", "inspect", image)
    volumes = [owner + "-" + suffix for suffix in ("files", "home", "env")]
    scratch = planning.scratch_name(owner, binding)
    try:
        for volume in volumes:
            docker("volume", "create", "--label", "workbench.test=" + prefix, volume)
        docker(
            "create",
            "--name",
            owner,
            "--network",
            "none",
            "--user",
            "0",
            "-v",
            volumes[0] + ":/workspace",
            "-v",
            volumes[1] + ":/home/dify",
            "-v",
            volumes[2] + ":/opt/user-env",
            "-e",
            "SHELLCTL_AUTH_TOKEN=test-secret",
            "--entrypoint",
            "/bin/sleep",
            image,
            "600",
        )
        docker("start", owner)
        setup = (
            "from pathlib import Path; import os; "
            f"Path('/workspace/conversations/{binding}').mkdir(parents=True); "
            "Path('/opt/workbench-global').mkdir(parents=True,exist_ok=True); "
            "paths=['/workspace/source.txt','/home/dify/personal.txt','/opt/user-env/installed.txt',"
            "'/opt/workbench-global/SKILL.md']; "
            "[(Path(p).write_text('source evidence'), os.chmod(p,0o644)) for p in paths]"
        )
        docker("exec", owner, "/usr/local/bin/python", "-I", "-S", "-c", setup)

        def inspect(script, **options):
            return planning.inspect_plan(
                key,
                payload(
                    binding_id=binding,
                    script=script,
                    admission_ticket=admission.ticket(key, binding),
                    **options,
                ),
                docker=docker,
                ensure=lambda value: None,
                identity=lambda value: (owner, "test-secret"),
                image=image,
                prefix=prefix,
                manager_python=("/usr/local/bin/python", "-I", "-S"),
                admission=admission,
            )

        inspect.stop = lambda: planning.stop_inspections(
            key, binding, docker=docker, prefix=prefix, admission=admission
        )
        inspect.docker = docker
        inspect.owner, inspect.key, inspect.binding, inspect.prefix = (
            owner,
            key,
            binding,
            prefix,
        )
        yield inspect
    finally:
        planning.stop_inspections(
            key, binding, docker=docker, prefix=prefix, admission=admission
        )
        docker("rm", "-f", owner, check=False)
        for volume in [*volumes, scratch]:
            # Exact names were created only by this fixture; no user volumes.
            docker("volume", "rm", volume, check=False)


def test_real_readonly_resources_offline_network_temporary_work_and_images(
    real_inspection,
):
    script = r"""python - <<'PY'
import base64, json, os, socket
from pathlib import Path
checks = {}
for path in ('/workspace/source.txt', '/home/dify/personal.txt', '/opt/user-env/installed.txt', '/opt/workbench-global/SKILL.md'):
    checks[path] = Path(path).read_text() == 'source evidence'
    try:
        Path(path).write_text('modified')
        checks[path + ':readonly'] = False
    except OSError:
        checks[path + ':readonly'] = True
try:
    Path('/workspace/created.txt').write_text('deliverable')
    checks['workspace-create-denied'] = False
except OSError:
    checks['workspace-create-denied'] = True
try:
    socket.create_connection(('1.1.1.1', 443), timeout=1).close()
    checks['offline'] = False
except OSError:
    checks['offline'] = True
checks['no-credential'] = 'SHELLCTL_AUTH_TOKEN' not in os.environ
Path('/tmp/analysis.txt').write_text('temporary result')
Path('/tmp/preview.png').write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+a1ioAAAAASUVORK5CYII='))
print(json.dumps(checks))
PY"""
    result = real_inspection(script, preview_paths=["/tmp/preview.png"])
    assert result["exit_code"] == 0, result
    assert all(json.loads(result["output"]).values()), result
    assert result["previews"][0]["media_type"] == "image/png"
    again = real_inspection("cat /tmp/analysis.txt")
    assert again["output"] == "temporary result"


def test_real_timeouts_output_bounds_and_background_cleanup(real_inspection):
    timeout = real_inspection("sleep 10", timeout=1)
    assert timeout["timed_out"] and timeout["exit_code"] == 124
    large = real_inspection("python -c \"print('a' * (11 * 1024 * 1024))\"")
    assert large["output_truncated"] and len(large["output"]) < 34000
    assert real_inspection(f"wc -c < {large['output_path']}")["output"].strip() == str(
        10 * 1024 * 1024
    )
    real_inspection(
        "(sleep 2; echo leaked > /tmp/background.txt) >/tmp/background.log 2>&1 &"
    )
    time.sleep(2.5)
    assert real_inspection("test ! -e /tmp/background.txt")["exit_code"] == 0


@pytest.mark.skipif(
    not os.environ.get("WORKBENCH_PLAN_TEST_OFFICE"),
    reason="Requires the sandbox image's installed Office parsing environment",
)
def test_real_document_parsing_and_pdf_preview(real_inspection):
    # Fixtures and rendered analysis stay in scratch, as real inspection does.
    result = real_inspection(
        r"""/opt/office/python/bin/python - <<'PY'
import json
from pathlib import Path
from docx import Document
import pymupdf
import pypdfium2

document = Document()
document.add_paragraph('Source amount: 123.45')
document.save('/tmp/source.docx')
pdf = pymupdf.open()
page = pdf.new_page()
page.insert_text((72, 72), 'Source amount: 123.45')
pdf.save('/tmp/source.pdf')
pdf.close()
with pymupdf.open('/tmp/source.pdf') as parsed:
    pdf_text = parsed[0].get_text()
word_text = Document('/tmp/source.docx').paragraphs[0].text
rendered = pypdfium2.PdfDocument('/tmp/source.pdf')
rendered[0].render(scale=1).to_pil().save('/tmp/document-preview.png')
print(json.dumps({'pdf': pdf_text.strip(), 'word': word_text,
                  'preview_bytes': Path('/tmp/document-preview.png').stat().st_size}))
PY""",
        timeout=20,
        preview_paths=["/tmp/document-preview.png"],
    )
    assert result["exit_code"] == 0, result
    observed = json.loads(result["output"])
    assert observed["pdf"] == observed["word"] == "Source amount: 123.45"
    assert observed["preview_bytes"] > 100
    assert result["previews"][0]["media_type"] == "image/png"
    assert result["warnings"] == []


def test_stop_inspections_scopes_containers_to_owner_and_binding(planning, admission):
    calls = []
    binding = str(uuid4())

    def docker(*args, **kwargs):
        calls.append(args)
        return SimpleNamespace(
            returncode=0, stdout="container-a\ncontainer-b\n" if args[0] == "ps" else ""
        )

    assert (
        planning.stop_inspections(
            "owner", binding, docker=docker, prefix="test", admission=admission
        )
        == 2
    )
    assert calls[0] == (
        "ps",
        "-aq",
        "--filter",
        "label=workbench=test",
        "--filter",
        "label=workbench.plan.owner=owner",
        "--filter",
        "label=workbench.plan.binding=" + binding,
    )
    assert calls[1:] == [("rm", "-f", "container-a"), ("rm", "-f", "container-b")]


def test_manager_stop_counts_inspections_when_owner_is_stopped(monkeypatch, tmp_path):
    directory = Path(__file__).parents[1] / "sandbox-manager"
    monkeypatch.syspath_prepend(str(directory))
    monkeypatch.setenv("WORKBENCH_SANDBOX_MANAGER_TOKEN", "test-token-" * 4)
    monkeypatch.setenv("WORKBENCH_MANAGER_STATE", str(tmp_path))
    spec = importlib.util.spec_from_file_location(
        "plan_stop_manager", directory / "server.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    calls = []

    def docker(*args, **kwargs):
        calls.append(args)
        return SimpleNamespace(
            returncode=1 if args[0] == "inspect" else 0,
            stdout="inspection-id" if args[0] == "ps" else "",
        )

    monkeypatch.setattr(module, "docker", docker)
    assert module.stop_binding(str(uuid4()), str(uuid4())) == {"stopped": 1}
    assert ("rm", "-f", "inspection-id") in calls


@pytest.mark.parametrize("stage", ["before_request", "during_initializer"])
def test_stop_revokes_delayed_and_preparing_inspections(planning, admission, stage):
    calls, containers = [], set()
    request = payload()
    binding = request["binding_id"]
    old_ticket = request["admission_ticket"] = admission.ticket("owner", binding)

    def docker(*args, **kwargs):
        calls.append(args)
        if args[0] == "create":
            containers.add(args[args.index("--name") + 1])
        elif args[0] == "ps":
            return SimpleNamespace(returncode=0, stdout="\n".join(containers))
        elif args[0] == "rm":
            containers.discard(args[-1])
        elif args[0] == "wait":
            assert stage == "during_initializer"
            planning.stop_inspections(
                "owner", binding, docker=docker, prefix="test", admission=admission
            )
        return SimpleNamespace(returncode=0, stdout="0")

    if stage == "before_request":
        planning.stop_inspections(
            "owner", binding, docker=docker, prefix="test", admission=admission
        )
    with pytest.raises(ValueError, match="调查执行已停止"):
        planning.inspect_plan(
            "owner",
            request,
            docker=docker,
            ensure=lambda _: None,
            identity=lambda _: ("test-owner", "token"),
            image="test-image",
            prefix="test",
            manager_python=("/usr/local/bin/python", "-I", "-S"),
            admission=admission,
        )
    assert not containers
    assert not any(call[0] == "exec" or "--volumes-from" in call for call in calls)
    # A manager restart cannot forget a completed stop or reopen its old ticket.
    restarted = planning.InspectionAdmission(admission.state, admission.lock)
    with pytest.raises(ValueError, match="调查执行已停止"):
        with restarted.guard("owner", binding, old_ticket):
            pytest.fail("old execution was reopened")
    with restarted.guard("owner", binding, restarted.ticket("owner", binding)):
        pass


def test_stop_does_not_confirm_failed_container_cleanup(planning, admission):
    def docker(*args, **kwargs):
        return SimpleNamespace(
            returncode=1 if args[0] == "rm" else 0,
            stdout="still-running" if args[0] == "ps" else "",
        )

    with pytest.raises(RuntimeError, match="无法确认"):
        planning.stop_inspections(
            "owner", str(uuid4()), docker=docker, prefix="test", admission=admission
        )


def test_real_stop_kills_active_inspection_before_returning(real_inspection):
    from concurrent.futures import ThreadPoolExecutor

    inspect = real_inspection
    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(
            inspect,
            "touch /tmp/stop-ready; sleep 30; touch /tmp/after-stop",
            timeout=40,
        )
        deadline = time.monotonic() + 15
        ready = False
        while time.monotonic() < deadline and not future.done():
            result = inspect.docker(
                "ps",
                "-q",
                "--filter",
                "label=workbench=" + inspect.prefix,
                "--filter",
                "label=workbench.plan.binding=" + inspect.binding,
            )
            for name in result.stdout.split():
                if not inspect.docker(
                    "exec", name, "test", "-f", "/tmp/stop-ready", check=False
                ).returncode:
                    ready = True
                    break
            if ready:
                break
            time.sleep(0.05)
        assert ready, "inspection did not reach its running script"
        assert inspect.stop() >= 1
        with pytest.raises((ValueError, subprocess.CalledProcessError)):
            future.result(timeout=10)
    remaining = inspect.docker(
        "ps",
        "-aq",
        "--filter",
        "label=workbench=" + inspect.prefix,
        "--filter",
        "label=workbench.plan.binding=" + inspect.binding,
    )
    assert remaining.stdout.strip() == ""
    assert (
        inspect("test -f /tmp/stop-ready && test ! -f /tmp/after-stop")["exit_code"]
        == 0
    )
