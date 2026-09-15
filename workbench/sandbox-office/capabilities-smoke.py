"""Generate representative office artifacts in the caller's writable directory."""
import importlib
import json
from pathlib import Path
import subprocess
import sys

out = Path(sys.argv[1]).resolve()
out.mkdir(parents=True, exist_ok=True)
result = {}

for module in ('docxcompose', 'plotly', 'kaleido', 'altair', 'vl_convert', 'networkx',
               'graphviz', 'pydot', 'cairosvg', 'svglib', 'weasyprint', 'numpy_financial',
               'cn2an', 'striprtf', 'magic', 'extract_msg', 'html5lib', 'cssselect'):
    importlib.import_module(module)
result['new_python_imports'] = 18

from docx import Document
from docxcompose.composer import Composer
from docxtpl import DocxTemplate
from openpyxl import load_workbook
from PIL import Image
import altair as alt
import cairosvg
import cn2an
import graphviz
import networkx as nx
import numpy_financial as npf
import pandas as pd
import plotly.graph_objects as go
import pydot
import pymupdf
from reportlab.graphics import renderPDF
from striprtf.striprtf import rtf_to_text
from svglib.svglib import svg2rlg
from weasyprint import HTML
import xlsxwriter

for index, text in enumerate(('材料目录', '证据说明')):
    document = Document()
    document.add_heading(text, 0)
    document.add_paragraph('中文文档处理测试')
    document.save(out / f'part-{index}.docx')
composer = Composer(Document(out / 'part-0.docx'))
composer.append(Document(out / 'part-1.docx'))
composer.save(out / 'merged.docx')
merged = '\n'.join(p.text for p in Document(out / 'merged.docx').paragraphs)
assert '材料目录' in merged and '证据说明' in merged
template_doc = Document()
template_doc.add_paragraph('金额：{{ amount }}；大写：{{ upper }}')
template_doc.save(out / 'template.docx')
template = DocxTemplate(out / 'template.docx')
template.render({'amount': '1234.56', 'upper': cn2an.an2cn('1234.56', 'rmb')})
template.save(out / 'filled.docx')
assert '壹仟贰佰叁拾肆元伍角陆分' in Document(out / 'filled.docx').paragraphs[0].text
assert rtf_to_text(r'{\rtf1\ansi Sample \b evidence\b0}') == 'Sample evidence'
result['document_merge_template_rtf'] = 'passed'

book = xlsxwriter.Workbook(out / 'amounts.xlsx')
sheet = book.add_worksheet('款项')
money = book.add_format({'num_format': '#,##0.00'})
sheet.write_row('A1', ['事项', '金额'])
sheet.write_row('A2', ['第一笔', 1200])
sheet.write_row('A3', ['第二笔', 3400])
sheet.write_formula('B4', '=SUM(B2:B3)', money, 4600)
sheet.add_table('A1:B3', {'columns': [{'header': '事项'}, {'header': '金额'}]})
sheet.set_column('A:A', 24)
sheet.set_column('B:B', 18, money)
chart = book.add_chart({'type': 'column'})
chart.add_series({'categories': '=款项!$A$2:$A$3', 'values': '=款项!$B$2:$B$3'})
sheet.insert_chart('D2', chart)
book.close()
frame = pd.read_excel(out / 'amounts.xlsx')
assert frame.iloc[:2]['金额'].sum() == 4600
assert round(float(npf.pv(0, 12, -100)), 2) == 1200
subprocess.run([sys.executable, '/opt/office/office.py', 'recalc', str(out / 'amounts.xlsx'),
                str(out / 'recalculated.xlsx')], check=True, timeout=120, capture_output=True, text=True)
assert load_workbook(out / 'recalculated.xlsx', data_only=True)['款项']['B4'].value == 4600
result['spreadsheet_chart_recalc_financial'] = 'passed'

figure = go.Figure(go.Bar(x=['一月', '二月', '三月'], y=[12, 18, 25]))
figure.update_layout(title='款项统计', font_family='Noto Sans CJK SC', width=700, height=420)
for suffix in ('png', 'svg', 'pdf'):
    figure.write_image(out / ('plotly.' + suffix))
figure.write_html(out / 'plotly.html', include_plotlyjs=True)
chart = alt.Chart(pd.DataFrame({'月份': ['一月', '二月'], '金额': [1200, 3400]})).mark_bar().encode(x='月份:N', y='金额:Q')
chart.save(out / 'altair.svg')
chart.save(out / 'altair.png')
chart.save(out / 'altair.html', inline=True)
relation = nx.DiGraph([('申请人', '被申请人'), ('被申请人', '担保人')])
assert nx.has_path(relation, '申请人', '担保人')
dot = graphviz.Digraph(graph_attr={'fontname': 'Noto Sans CJK SC'}, node_attr={'fontname': 'Noto Sans CJK SC'})
for edge in relation.edges:
    dot.edge(*edge)
(out / 'relations.svg').write_bytes(dot.pipe(format='svg'))
assert pydot.graph_from_dot_data(dot.source)
result['plotly_altair_graphviz'] = 'passed'

svg = '<svg xmlns="http://www.w3.org/2000/svg" width="320" height="180"><rect width="320" height="180" fill="#e0f2fe"/><circle cx="160" cy="85" r="45" fill="#2563eb"/><text x="140" y="160">Chart</text></svg>'
(out / 'vector.svg').write_text(svg)
cairosvg.svg2png(bytestring=svg.encode(), write_to=str(out / 'vector.png'))
cairosvg.svg2pdf(bytestring=svg.encode(), write_to=str(out / 'vector.pdf'))
drawing = svg2rlg(str(out / 'vector.svg'))
renderPDF.drawToFile(drawing, str(out / 'reportlab-vector.pdf'))
subprocess.run(['rsvg-convert', '-o', str(out / 'rsvg.png'), str(out / 'vector.svg')], check=True, timeout=30)
subprocess.run(['convert', str(out / 'vector.png'), '-resize', '160x90', str(out / 'thumbnail.png')], check=True, timeout=30)
result['svg_vector_and_raster'] = 'passed'

html = '''<!doctype html><meta charset="utf-8"><style>
@page {size:A4;margin:20mm;@bottom-center{content:counter(page);}}
body{font-family:"Noto Sans CJK SC";}h1{color:#1d4ed8;}section+section{break-before:page;}
table{border-collapse:collapse;width:100%;}td,th{border:1px solid #777;padding:8px;}
</style><section><h1>材料目录</h1><table><tr><th>事项</th><th>金额</th></tr><tr><td>测试款项</td><td>4600</td></tr></table></section><section><h1>证据说明</h1><p>中文分页打印测试。</p></section>'''
(out / 'print.html').write_text(html, encoding='utf8')
HTML(string=html, base_url=str(out)).write_pdf(out / 'print.pdf')
with pymupdf.open(out / 'print.pdf') as document:
    assert len(document) == 2
    assert '材料目录' in document[0].get_text() and '证据说明' in document[1].get_text()
    document[0].get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(out / 'print.png')
result['html_chinese_paged_pdf'] = 'passed'
subprocess.run(['sqlite3', str(out / 'table.sqlite'), 'CREATE TABLE amounts (amount INTEGER); INSERT INTO amounts VALUES (4600);'], check=True)
assert subprocess.check_output(['sqlite3', str(out / 'table.sqlite'), 'SELECT SUM(amount) FROM amounts;'], text=True).strip() == '4600'
result['sqlite_table'] = 'passed'

subprocess.run(['node', '/opt/office/capabilities-smoke.cjs', str(out)], check=True, timeout=180)
result['node'] = json.loads((out / 'node-capabilities.json').read_text())
for path in out.glob('*.png'):
    with Image.open(path) as image:
        image.verify()
for path in out.glob('*.pdf'):
    with pymupdf.open(path) as document:
        assert len(document) > 0
result.update(status='passed', artifacts=len([p for p in out.rglob('*') if p.is_file()]))
(out / 'capabilities-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False), flush=True)
