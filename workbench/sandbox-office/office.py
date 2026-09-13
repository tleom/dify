"""Office conversion, verified recalculation and image previews for the user sandbox."""
import argparse
import json
import os
from pathlib import Path
import signal
import subprocess
import sys


def transform(source, target, action='convert', timeout=120):
    worker = Path(__file__).with_name('uno_worker.py')
    process = subprocess.Popen(['/usr/bin/python3', str(worker), str(source), str(target), action],
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding='utf-8',
                               start_new_session=True)
    try:
        stdout, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        os.killpg(process.pid, signal.SIGTERM)
        try:
            process.communicate(timeout=3)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid, signal.SIGKILL)
            process.communicate()
        raise TimeoutError('LibreOffice timed out; output is not verified') from None
    if process.returncode:
        raise RuntimeError(stdout.strip() or stderr[-2000:] or 'Office operation failed')
    return json.loads(stdout)


def preview(source, directory, dpi=120):
    import pymupdf

    source = Path(source).resolve()
    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    pdf = source if source.suffix.lower() == '.pdf' else directory / (source.stem + '.pdf')
    if pdf != source:
        transform(source, pdf)
    images = []
    with pymupdf.open(pdf) as document:
        if document.needs_pass:
            raise ValueError('The PDF requires a password before previewing')
        for index, page in enumerate(document):
            image = directory / f'page-{index + 1:04}.png'
            if image.exists():
                raise FileExistsError(image)
            page.get_pixmap(matrix=pymupdf.Matrix(dpi / 72, dpi / 72)).save(image)
            images.append({'page': index + 1, 'image': str(image)})
    return {'status': 'success', 'pdf': str(pdf), 'page_count': len(images), 'pages': images}


def check_formulas(path):
    from openpyxl import load_workbook

    formulas = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    errors = {}
    missing = []
    count = 0
    try:
        for sheet in formulas:
            for row in sheet:
                for cell in row:
                    value = cached[sheet.title][cell.coordinate]
                    location = f'{sheet.title}!{cell.coordinate}'
                    if cell.data_type == 'f':
                        count += 1
                        # LibreOffice stores a valid empty-string result as t="str" with an empty v.
                        if value.value is None and value.data_type != 'str':
                            missing.append(location)
                    if value.data_type == 'e':
                        errors.setdefault(str(value.value), []).append(location)
    finally:
        formulas.close()
        cached.close()
    return {
        'status': 'errors_found' if errors or missing else 'success', 'total_formulas': count,
        'total_errors': sum(len(items) for items in errors.values()),
        'error_summary': {key: {'count': len(items), 'locations': items} for key, items in errors.items()},
        'missing_cached_values': missing,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['convert', 'preview', 'recalc', 'accept-changes'])
    parser.add_argument('input')
    parser.add_argument('output', help='Output file, or a new preview directory')
    parser.add_argument('--dpi', type=int, default=120)
    parser.add_argument('--timeout', type=int, default=120)
    args = parser.parse_args()
    try:
        result = preview(args.input, args.output, args.dpi) if args.action == 'preview' else transform(args.input, args.output, args.action, args.timeout)
        if args.action == 'recalc':
            result.update(check_formulas(args.output))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result['status'] == 'success' else 1
    except Exception as error:
        print(json.dumps({'status': 'failed', 'error': str(error)}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
