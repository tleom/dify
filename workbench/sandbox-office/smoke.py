"""Exercise document generation, Office preview, Chinese text and both browser APIs."""
import importlib
import json
from pathlib import Path
import subprocess
import sys
import uuid

out = Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/office-smoke')
out.mkdir(parents=True, exist_ok=True)
for module in ('docx', 'docxtpl', 'pptx', 'openpyxl', 'xlsxwriter', 'xlrd', 'pandas', 'numpy', 'scipy', 'matplotlib', 'seaborn', 'PIL', 'pymupdf', 'pypdf', 'pdfplumber', 'pdf2image', 'reportlab', 'pytesseract', 'markitdown', 'lxml', 'defusedxml', 'bs4', 'requests', 'httpx', 'yaml', 'playwright'):
    importlib.import_module(module)

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook
from markitdown import MarkItDown
from playwright.sync_api import sync_playwright
import pymupdf

text = '中文预览验证'
doc = Document()
run = doc.add_paragraph().add_run(text)
run.font.name = 'Noto Sans CJK SC'
fonts = OxmlElement('w:rFonts')
fonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
run._element.get_or_add_rPr().append(fonts)
doc.save(out / 'document.docx')
book = Workbook()
sheet = book.active
sheet.append([text, '金额'])
sheet.append(['甲', 10])
sheet.append(['乙', 20])
sheet.append(['合计', '=SUM(B2:B3)'])
sheet.column_dimensions['A'].width = 28
book.save(out / 'sheet.xlsx')
subprocess.run(['node', '/opt/office/smoke.cjs', str(out)], check=True, timeout=90)

result = {'ok': True, 'previews': {}, 'python_browser': False}
pdfs = out / 'preview'
pdfs.mkdir(exist_ok=True)
for name in ('document.docx', 'sheet.xlsx', 'slides.pptx'):
    profile = f'file:///tmp/lo-smoke-{uuid.uuid4().hex}'
    subprocess.run(['soffice', f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf', '--outdir', str(pdfs), str(out / name)], check=True, timeout=90, capture_output=True, text=True)
    pdf = pdfs / (Path(name).stem + '.pdf')
    assert pdf.is_file() and pdf.stat().st_size > 1000, name
    with pymupdf.open(pdf) as pages:
        contents = ''.join(page.get_text() for page in pages)
        assert text in contents, (name, contents)
        pages[0].get_pixmap(matrix=pymupdf.Matrix(1.2, 1.2)).save(pdfs / (Path(name).stem + '.png'))
        result['previews'][name] = {'pages': len(pages), 'chinese_text': True}
    extracted = MarkItDown().convert(out / name).text_content
    assert text in extracted, (name, extracted)

subprocess.run(['pdftoppm', '-f', '1', '-singlefile', '-png', '-r', '72', str(pdfs / 'document.pdf'), str(out / 'poppler-preview')], check=True, timeout=30, capture_output=True)
assert (out / 'poppler-preview.png').stat().st_size > 1000
with sync_playwright() as p:
    browser = p.chromium.launch(channel='chrome', headless=True, args=['--no-sandbox', '--disable-dev-shm-usage'])
    try:
        page = browser.new_page()
        page.set_content(f'<title>{text}</title><h1>{text}</h1>')
        assert page.title() == text
        page.screenshot(path=str(out / 'chrome-python.png'))
        result['python_browser'] = True
    finally:
        browser.close()
result['chrome'] = json.loads((out / 'node-result.json').read_text())
(out / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps(result, ensure_ascii=False))
